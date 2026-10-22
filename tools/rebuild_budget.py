#!/usr/bin/env python3
"""预算表重建脚本：输入只有 A-H（院系提供）的预算表，重建 I-S 呈现列与 T-AC 辅助列。

用法：
    python3 rebuild_budget.py 输入.xlsx [-o 输出.xlsx] [--sheet 表名] [--fares fares.json]

列布局（与 2026 暑假预算表 0626 一致）：
    A-H   院系提供：序号/实践地点/拟主要参访点位/承办院系/拟参访日期/支队人数/(空)/院系已有支持金额
    I-S   主要呈现列（公式）：实践地点/参访天数/支队人数/往返交通费+合计/住宿+合计/保险+合计/总计/备注
    T-AC  计算辅助列：单人交通费(T，唯一外部数据) … 队伍总额/报销比例（校研团不重复支持，已移除）

T 列由本脚本按 B 列实践地解析城市 → 展开行程链（北京→A→B→北京）→ 查 fares.json 写入；
查不到的段标记"待查: XX-XX"，需人工补票价（补进 fares.json 后重跑即可）。
加 --online 时，缺失分段自动调用 fare_query.py 联网查 12306 实时最高二等座价并写回 fares.json；
加 --refresh 时（每期做预算前建议用），先把本表涉及的全部分段联网重查最新价再重建。
"""
import argparse, json, re, sys
from datetime import date
import openpyxl

KNOWN_CITIES = ["上海", "西安", "武汉", "合肥", "广州", "深圳", "珠海", "成都",
                "福州", "宁德", "南京", "苏州", "无锡", "杭州", "重庆", "江门",
                "青岛", "沈阳", "大连", "长沙", "郑州", "天津"]

def known_cities(fares):
    """已知城市 = 内置基础表 ∪ fares.json 分段端点（除北京）。新增分段后自动可识别。"""
    cities = set(KNOWN_CITIES)
    for key in fares.get("segments", {}):
        for ep in key.split("-"):
            if ep != "北京":
                cities.add(ep)
    return sorted(cities, key=len, reverse=True)  # 长名优先（如"石家庄"先于"石家"误配）

def parse_cities(text, cities=None):
    """从实践地字符串中按出现顺序提取已知城市。"""
    t = re.sub(r"\s+", "", str(text or ""))
    found = sorted(((t.find(c), c) for c in (cities or KNOWN_CITIES) if c != "北京" and c in t))
    return [c for _, c in found if _ >= 0]

def parse_days(text, year=2026):
    """解析 '7.8-7.10' / '7.28-8.2' / '12.9-12.13' → 天数（含首尾）。"""
    rng = parse_date_range(text, date(year, 1, 1))
    return (rng[1] - rng[0]).days + 1 if rng else None

def parse_date_range(text, today):
    """解析 '7.19-7.23' → (起始日, 结束日)。年份取今年，已过去则取明年。"""
    m = re.search(r"(\d+)\.(\d+)\s*-\s*(?:(\d+)\.)?(\d+)", str(text or ""))
    if not m:
        return None
    m1, d1 = int(m.group(1)), int(m.group(2))
    m2, d2 = int(m.group(3) or m.group(1)), int(m.group(4))
    for yr in (today.year, today.year + 1):
        try:
            s, e = date(yr, m1, d1), date(yr, m2, d2)
        except ValueError:
            continue
        if e >= today:
            return s, e
    return None

def fmt_num(x):
    return str(int(x)) if float(x) == int(x) else str(x)

def fare_for(cities, fares):
    """返回 (T列写入值, 状态说明)。segments 优先，其次 unsplit 总价 / roundtrip。"""
    segs, rt, unsplit = fares["segments"], fares["roundtrip"], fares["unsplit"]
    if not cities:
        return "待查: 未识别实践地", "未识别城市"
    chain = ["北京"] + cities + ["北京"]
    prices, missing = [], []
    for i in range(len(chain) - 1):
        key = f"{chain[i]}-{chain[i+1]}"
        rev = f"{chain[i+1]}-{chain[i]}"
        if key in segs:
            prices.append(segs[key])
        elif rev in segs:  # 高铁往返近似同价，反向兜底
            prices.append(segs[rev])
        else:
            missing.append(key)
    if not missing:
        return "=" + "+".join(fmt_num(p) for p in prices), "ok(分段)"
    if len(cities) == 1 and cities[0] in rt:
        return rt[cities[0]], "ok(往返总价)"
    if len(cities) > 1 and "、".join(cities) in unsplit:
        return unsplit["、".join(cities)], "ok(多城总价, 分段待补)"
    return "待查: " + "，".join(missing), "缺票价: " + "，".join(missing)

def fare_basis(cities, fares):
    """AD列「交通费取价依据」：逐段写明票价、来源（12306实时车次+查询日期 / 缓存）。"""
    segs, rt, unsplit = fares["segments"], fares["roundtrip"], fares["unsplit"]
    meta = fares.get("meta", {})
    if not cities:
        return "未识别实践地，待人工确认"

    def seg_text(key, rev):
        k = key if key in segs else (rev if rev in segs else None)
        if k is None:
            return f"{key} 待查"
        m = meta.get(k, {})
        if m.get("manual"):
            src = "人工录入·需复核" + (f"·{m['note']}" if m.get("note") else "")
        elif m.get("train"):
            if m.get("window"):
                w0, w1 = m["window"].split("~")
                src = f"12306实时·{m['train']}·窗口{w0[5:]}~{w1[5:]}·最高日{m['query_date'][5:]}"
            else:
                src = f"12306实时·{m['train']}·{m['query_date']}查"
        else:
            src = "缓存·源自2026暑假预算表"
        return f"{key} {fmt_num(segs[k])}元[{src}]"

    chain = ["北京"] + cities + ["北京"]
    parts = [seg_text(f"{chain[i]}-{chain[i+1]}", f"{chain[i+1]}-{chain[i]}")
             for i in range(len(chain) - 1)]
    if all("待查" not in p for p in parts):
        return "；".join(parts)
    if len(cities) == 1 and cities[0] in rt:
        return f"{cities[0]}往返总价 {fmt_num(rt[cities[0]])}元[缓存·源自2026暑假预算表]"
    if len(cities) > 1 and "、".join(cities) in unsplit:
        return ("、".join(cities) + f"总价 {fmt_num(unsplit['、'.join(cities)])}元"
                "[缓存·源自2026暑假预算表，分段待补]；" + "；".join(parts))
    return "；".join(parts)

def row_formulas(r):
    """I-S 呈现列 + U-AC 辅助列公式（T/V/Z 除外）。

    保险统一口径（据「审核报销」skill 就协规则库 career-association.json）：
    ≤30 元/人/周，AA = Z × 人数 × ceil(天数/7) × 比例；≤7 天即 30 元/人。
    """
    return {
        "I": f"=B{r}",
        "K": f"=F{r}",
        "L": f'=T{r}&"/人"&"x"&TEXT(AC{r},"# ?/?")',
        "M": f"=U{r}",
        "N": f'=V{r}/2&"/人·天*"&TRIM(TEXT(AC{r},"# ?/?"))&"*"&X{r}&"晚"',
        "O": f"=Y{r}",
        "P": f'=Z{r}&"/人·周*"&CEILING(J{r}/7,1)&"周"&"x"&TEXT(AC{r},"# ?/?")',
        "Q": f"=AA{r}",
        "R": f"=AB{r}",
        "S": f'=IF(OR(ISNUMBER(SEARCH("就协",D{r})),ISNUMBER(SEARCH("就业协会",D{r}))),'
             f'"就协全部","院系"&TEXT(1-AC{r},"# ?/?")&"，剩余就协")',
        "U": f"=K{r}*T{r}*AC{r}",
        "W": f"=CEILING(K{r}/2,1)",
        "X": f"=J{r}-1",
        "Y": f"=V{r}*W{r}*X{r}*AC{r}",
        "AA": f"=Z{r}*K{r}*CEILING(J{r}/7,1)*AC{r}",
        "AB": f"=U{r}+Y{r}+AA{r}",
        "AC": f'=IF(OR(ISNUMBER(SEARCH("就协",D{r})),ISNUMBER(SEARCH("就业协会",D{r}))),1,2/3)',
    }

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("-o", "--output")
    ap.add_argument("--sheet", default="各支队预算按实际人数校准")
    ap.add_argument("--fares", default="fares.json")
    ap.add_argument("--online", action="store_true",
                    help="缺失分段联网查 12306 最高价（fare_query.py）并写回 fares.json")
    ap.add_argument("--refresh", action="store_true",
                    help="预算前预处理：本表涉及的全部分段联网重查最新最高价（覆盖缓存）后再重建")
    ap.add_argument("--refresh-days", type=int, default=14,
                    help="--refresh 的查询窗口天数（默认14=预售期全量，覆盖全周浮动票价）")
    ap.add_argument("--start-row", type=int, default=7)
    args = ap.parse_args()

    fares = json.load(open(args.fares, encoding="utf-8"))
    cities_list = known_cities(fares)
    wb = openpyxl.load_workbook(args.input)
    ws = wb[args.sheet] if args.sheet in wb.sheetnames else wb.active
    out = args.output or args.input.replace(".xlsx", "_重建.xlsx")

    # AD 列：交通费取价依据（逐段来源留痕）
    ws[f"AD{args.start_row - 1}"] = "交通费取价依据"
    ws.column_dimensions["AD"].width = 70
    from openpyxl.styles import Alignment, Font
    ws[f"AD{args.start_row - 1}"].font = Font(bold=True)

    fare_client = None
    fares_dirty = False

    def fetch_missing(cities):
        """--online：把行程链中 fares.json 缺失的分段联网补齐。"""
        nonlocal fare_client, fares_dirty
        chain = ["北京"] + cities + ["北京"]
        segs = fares["segments"]
        for i in range(len(chain) - 1):
            key = f"{chain[i]}-{chain[i+1]}"
            rev = f"{chain[i+1]}-{chain[i]}"
            if key in segs or rev in segs:
                continue
            if fare_client is None:
                from fare_query import FareClient
                fare_client = FareClient()
            try:
                r = fare_client.max_fare(chain[i], chain[i + 1])
                segs[key] = r["price"]
                fares.setdefault("meta", {})[key] = {
                    "query_date": r["query_date"], "train": r["train"]}
                fares_dirty = True
                print(f"  联网补价 {key} = {r['price']} [{r['train']} @{r['query_date']}]")
            except Exception as e:
                print(f"  联网查价失败 {key}: {e}", file=sys.stderr)

    def refresh_needed():
        """--refresh：预算前把本表所有行程链分段按窗口联网重查（不论缓存是否已有）。

        窗口选择：若表内拟参访日期与 12306 预售期（约15天）有交集，按交集窗口查；
        否则默认查"明天起 --refresh-days 天"。窗口内逐日取最高价，覆盖浮动票价。
        """
        nonlocal fare_client, fares_dirty
        from datetime import timedelta
        today = date.today()
        sale_start, sale_end = today + timedelta(days=1), today + timedelta(days=14)
        needed, windows = [], []
        rr = args.start_row
        while ws[f"A{rr}"].value is not None or ws[f"B{rr}"].value is not None:
            cities = parse_cities(ws[f"B{rr}"].value, cities_list)
            chain = ["北京"] + cities + ["北京"] if cities else []
            for i in range(len(chain) - 1):
                key = f"{chain[i]}-{chain[i+1]}"
                rev = f"{chain[i+1]}-{chain[i]}"
                if key not in needed and rev not in needed:
                    needed.append(key)
            rng = parse_date_range(ws[f"E{rr}"].value, today)
            if rng:  # 与预售期取交集
                s, e = max(rng[0], sale_start), min(rng[1], sale_end)
                if s <= e:
                    windows.append((s, e))
            rr += 1
        if not needed:
            return
        if windows:  # 实践日期落在预售期内 → 按拟定窗口查
            win_start, win_end = min(s for s, _ in windows), max(e for _, e in windows)
        else:        # 否则默认"明天起 N 天"
            win_start = sale_start
            win_end = min(sale_start + timedelta(days=args.refresh_days - 1), sale_end)
        days = (win_end - win_start).days + 1
        from fare_query import FareClient, _meta_of, _fmt_tag
        fare_client = FareClient()
        print(f"预算前刷新票价：{len(needed)} 个分段 × 窗口 {win_start}~{win_end}"
              f"（{days}天，{'按表内拟定日期' if windows else '默认预售期窗口'}）")
        for key in needed:
            frm, to = key.split("-", 1)
            try:
                r = fare_client.max_fare_range(frm, to, start=win_start, days=days)
                old = fares["segments"].get(key)
                was_manual = fares.setdefault("meta", {}).get(key, {}).get("manual")
                fares["segments"][key] = r["price"]
                fares["meta"][key] = _meta_of(r)
                fares_dirty = True
                diff = "" if old is None else ("（新增）" if not isinstance(old, (int, float))
                     else f"（原人工录入 {old}，现已有直达实时价）" if was_manual
                     else f"（原 {old}）" if abs(r["price"] - old) >= 0.01 else "（不变）")
                print(f"  {key} = {r['price']} [{_fmt_tag(r)}]{diff}")
            except Exception as e:
                if fares.setdefault("meta", {}).get(key, {}).get("manual"):
                    print(f"  {key} ⚑ 人工录入价 {fares['segments'].get(key)} "
                          f"仍无法联网复核（无直达车），请人工确认后使用", file=sys.stderr)
                else:
                    print(f"  {key} 刷新失败（保留缓存价）: {e}", file=sys.stderr)
        fares.setdefault("meta", {})["updated_at"] = str(date.today())

    if args.refresh:
        refresh_needed()

    report = []
    r = args.start_row
    while ws[f"A{r}"].value is not None or ws[f"B{r}"].value is not None:
        B, D, E = ws[f"B{r}"].value, ws[f"D{r}"].value, ws[f"E{r}"].value
        cities = parse_cities(B, cities_list)
        if args.online and not args.refresh and cities:
            fetch_missing(cities)
        t_val, status = fare_for(cities, fares)
        days = parse_days(E)

        for col, f in row_formulas(r).items():
            ws[f"{col}{r}"] = f
        ws[f"J{r}"] = days if days else "待填"
        ws[f"T{r}"] = t_val
        ad = ws[f"AD{r}"]
        ad.value = fare_basis(cities, fares)
        ad.alignment = Alignment(wrap_text=True, vertical="center")
        ad.font = Font(size=9, color="595959")
        # 实践地承担住宿时（A-H 任一文本含"承担住宿"），标间单价置 0
        row_text = "".join(str(ws[f"{c}{r}"].value or "") for c in "ABCDEFGH")
        lodging_free = "承担住宿" in row_text or "住宿自理" in row_text
        ws[f"V{r}"] = 0 if lodging_free else 300
        ws[f"Z{r}"] = 30
        if lodging_free:
            status += "；住宿由实践地承担(V=0)"

        report.append((r, str(B), cities, days, status))
        r += 1

    # 对账区：过会金额（手填）vs 预算合计（自动求和）vs 差额
    chk = r + 1
    from openpyxl.styles import Font as _Font
    ws[f"Q{chk}"] = "过会金额"
    ws[f"Q{chk+1}"] = "预算合计"
    ws[f"R{chk+1}"] = f"=SUM(R{args.start_row}:R{r - 1})"
    ws[f"Q{chk+2}"] = "差额"
    ws[f"R{chk+2}"] = f"=R{chk}-R{chk+1}"
    for rr in (chk, chk + 1, chk + 2):
        ws[f"Q{rr}"].font = _Font(bold=True)

    wb.save(out)
    if fares_dirty:
        json.dump(fares, open(args.fares, "w", encoding="utf-8"),
                  ensure_ascii=False, indent=2)
        print("fares.json 已同步更新（联网查询结果已写回）")
    print(f"已重建 {r - args.start_row} 行 → {out}\n")
    for row, b, cities, days, status in report:
        flag = "✓" if status.startswith("ok") else "✗"
        print(f"{flag} 行{row} {b!r:24} 城市={cities} 天数={days} [{status}]")
    if any(not s.startswith("ok") for _, _, _, _, s in report):
        print("\n存在待查票价，请补入 fares.json 的 segments 后重跑。", file=sys.stderr)

if __name__ == "__main__":
    main()
